"""流式导出引擎 —— 万条聊天不 OOM。

核心改进（对比旧版）：
1. message_stream() 流式读取：heapq.merge K 路归并，内存 O(分片数)
2. 增量写入 JSON/TXT/CSV/MD：直接写文件句柄，不构建巨型字符串
3. 并行媒体解密：multiprocessing.Pool 多进程 AES 解密
4. 双遍扫描：第一遍轻量采集元数据（计数/发送者/图片引用），第二遍流式写出
"""
import hashlib
import json
import os
import re
import shutil
import sqlite3
import time
from datetime import datetime
from pathlib import Path

from siwx import logger, media, sns_cdn, voice
from siwx import paths as _paths
from siwx.api_chat import (
    COMPOSITE_49_LABELS, FALLBACK_LABEL, _contact_names, has_refermsg,
)
from siwx.export_stream import (
    message_stream, stream_export_json, stream_export_txt,
    stream_export_csv, stream_export_md,
)

GENERATOR = "stories-in-wx"
EXPORT_VERSION = "1.0"

# 文本截断阈值（P3-8）。此前各格式互不一致（CSV 2000 / HTML 8000 / SQLite 8000
# / XLSX 30000）且 manifest 无记录；现统一到命名常量，并把每格式实际截断条数
# 写进 manifest["truncated"] 与 run_export 返回值。
# - HTML/SQLite 8000：html_template._msg_entry 对 rawContent 也按 8000 截断，
#   两处必须同步；超过 8000 字的单条消息实测极少。
# - XLSX 30000：Excel 单元格硬上限 32767 字符，留出余量。
# - CSV/JSON/TXT/Markdown/TOML 不截断（CSV 标准无单元格长度限制）。
HTML_RAW_CONTENT_LIMIT = 8000
SQLITE_RAW_CONTENT_LIMIT = 8000
XLSX_CELL_LIMIT = 30000

# 语音媒体在组合键 media_map 里的槽位前缀（与图片槽位隔离）
_VOICE_KEY = "\x00voice"


class ExportCancelled(Exception):
    """协作式取消：should_cancel() 返回 True 时在最近的检查点抛出。

    取消是优雅停止：已完成会话的产物保留，报告带 cancelled=True；
    与异常失败（error 项）严格区分。"""

# 兜底文案精确集合（P2 修复）：判定"content 恰好等于兜底标签原文"才算兜底。
# 此前用 startswith 前缀匹配——合法解析成功的链接卡是 "[链接] 标题"、
# 转账是 "[转账] 标题"、合并转发是 "[聊天记录] 标题（N 条…）"，全部被误计
# 为兜底，manifest 的 fallback_labels 严重虚高且逐条刷噪声日志。
_FALLBACK_LABELS_EXACT = frozenset(
    set(FALLBACK_LABEL.values()) | set(COMPOSITE_49_LABELS.values())
    | {"[链接]", "[引用]", "[聊天记录]", "[系统消息]"})
# "[类型N]"（N 任意数字）是未知类型的兜底文案，正则精确匹配
_FALLBACK_TYPE_RE = re.compile(r"^\[类型\d+\]$")


# Windows 保留设备名，不能作为文件/目录名
_WIN_RESERVED = {"CON", "PRN", "AUX", "NUL",
                 *(f"COM{i}" for i in range(1, 10)),
                 *(f"LPT{i}" for i in range(1, 10))}


def _safe_name(name: str) -> str:
    """把任意文本清洗为安全的文件/目录名。

    Bug 修复：原实现写作 `name = ch.replace(ch, "_")`，把 name 覆盖成了单个字符
    替换的结果（`ch` 只含一个字符，`ch.replace(ch, "_")` 恒等于 `"_"`），于是
    无论传入什么——包括联系人名称——都恒返回 `"_"`。这正是导出目录与文件名里
    会话名/联系人名永远为空的原因。
    """
    name = str(name or "")
    for ch in '<>:"/\\|?*':
        name = name.replace(ch, "_")
    for ch in "\r\n\t":
        name = name.replace(ch, " ")
    # Windows 不允许名称以点或空格结尾
    name = name[:48].strip().rstrip(".").strip()
    if not name:
        return "chat"
    if name.split(".")[0].upper() in _WIN_RESERVED:
        name = "_" + name
    return name


def collect_avatars(acc_out_dir: Path, usernames: list, dest: Path, progress=None):
    """从 head_image.db 提取头像 → avatars/<md5(username)>.jpg。

    审计 §4.1：写失败原为未处理异常，会让 run_export 的头像阶段整体崩溃；
    现逐条包 try + 计数，结束时汇总一条 detailed。

    P3-7：连接与查询同样包 `sqlite3.Error`——head_image.db 损坏/被占用时
    只降级为「没有头像」（返回已收集的 mapping），不再让整个会话导出失败。"""
    db = acc_out_dir / "head_image" / "head_image.db"
    mapping = {}
    if not db.is_file():
        logger.detailed("export", f"头像库不存在: {db.name}")
        return mapping
    conn = None
    hit = fail = 0
    try:
        conn = sqlite3.connect(db)
        for i, un in enumerate(usernames):
            row = conn.execute(
                "SELECT image_buffer FROM head_image WHERE username=?", (un,)).fetchone()
            if row and row[0]:
                fn = hashlib.md5(un.encode()).hexdigest() + ".jpg"
                try:
                    (dest / fn).write_bytes(row[0])
                    mapping[un] = f"avatars/{fn}"
                    hit += 1
                except OSError as e:
                    fail += 1
                    logger.detailed("export", f"头像写入失败 username={un}: {e}")
            if (i + 1) % 50 == 0 and progress:
                progress(0, f"头像 {i + 1}/{len(usernames)}")
    except sqlite3.Error as e:
        # 库损坏（file is not a database）/被锁：保留已收集的头像，继续导出
        logger.warn("export",
                    f"头像库读取失败，本次导出降级为无头像"
                    f"（已收集 {len(mapping)} 个，用户数 {len(usernames)}）: "
                    f"{db.name}: {e}")
    finally:
        if conn is not None:
            try:
                conn.close()
            except sqlite3.Error:
                pass
    logger.detailed("export",
                    f"头像提取: db_exists=True 用户数={len(usernames)} "
                    f"命中={hit} 写失败={fail}")
    return mapping


def _collect_metadata(acc_out_dir, chat, start_ts, end_ts, account, names=None):
    """第一遍：轻量扫描，只采集计数/发送者/媒体引用。内存 O(发送者数 + 媒体数)。"""
    senders = set()
    images = []       # (md5, bubble_md5, localId, ts) 引用
    voices = []       # (localId, serverId, ts) 引用
    stickers = []     # (md5, bubble_md5, localId, ts, cdnurl) 引用
    count = 0
    first_ts = last_ts = 0
    for msg in message_stream(acc_out_dir, chat, start_ts, end_ts, account, names):
        count += 1
        ts = msg["createTime"] or 0
        if count == 1:
            first_ts = ts
        last_ts = ts
        senders.add(msg["senderUsername"])
        if msg.get("md5") or msg.get("bubbleMd5"):
            images.append((msg.get("md5"), msg.get("bubbleMd5"),
                           msg["localId"], ts, chat))
        # 合并转发子项里的图片/表情（issue #31）：定位键用原始消息的
        # local_id/create_time，原始消息可能属于别的会话（跨会话转发），
        # attach 目录必须按 srcChat 算
        if msg.get("record"):
            for it in _iter_record_media(msg["record"]):
                images.append((it.get("md5"), it.get("bubbleMd5"),
                               it.get("srcLocalId") or 0,
                               it.get("srcCreateTime") or 0,
                               it.get("srcChat") or chat))
        if msg.get("localType") == 34:
            voices.append((msg["localId"], msg.get("platformMessageId") or "", ts))
        if msg.get("localType") == 47:
            url = (msg.get("sticker") or {}).get("url") or ""
            if url:
                stickers.append((msg.get("md5"), msg.get("bubbleMd5"),
                                 msg["localId"], ts, url))
    return count, first_ts, last_ts, senders, images, voices, stickers


def _media_key(md5, bubble_md5, local_id, ts) -> tuple:
    """图片媒体的组合键 (md5, localId, ts)。

    多分片会话各分片的 local_id 独立编号会撞号（真实库实测 19210 个撞号键 /
    350 个会话），旧版裸 localId 键会让跨分片同 localId 的两条图片消息互挂
    对方的图。md5 优先取消息 XML 提取值，缺失时用气泡 packed_info 里的
    md5；组合键下同 md5 同图共享是正确行为，不同 md5 不再互挂。
    """
    return ((md5 or bubble_md5 or ""), int(local_id or 0), int(ts or 0))


def _voice_key(local_id, ts) -> tuple:
    """语音媒体的组合键（语音无 md5，用固定槽位前缀隔离）。"""
    return (_VOICE_KEY, int(local_id or 0), int(ts or 0))


def _iter_record_media(record):
    """递归遍历合并转发子项里带媒体引用的条目（issue #31）。

    只产出有 md5/bubbleMd5 的子项（图片/表情）；嵌套的下一层记录继续下钻。
    子项的 srcLocalId/srcCreateTime 指向原始消息，是 attach/Bubble/Thumb
    三级来源的定位键。
    """
    for it in (record or {}).get("items") or []:
        if it.get("md5") or it.get("bubbleMd5"):
            yield it
        nested = it.get("record")
        if nested:
            yield from _iter_record_media(nested)


def _attach_record_media(record, media_map) -> None:
    """把已解密的媒体回填到合并转发子项（issue #31）。

    子项与顶层消息共用同一个 media_map：键是 (md5, 原始 local_id, ts)，
    与顶层消息的 local_id 天然不同，不会互挂。
    """
    for it in _iter_record_media(record):
        mf = media_map.get(_media_key(it.get("md5"), it.get("bubbleMd5"),
                                      it.get("srcLocalId") or 0,
                                      it.get("srcCreateTime") or 0))
        if mf and "/voice_" not in mf:
            it["mediaFile"] = mf


def _attach_media(msg: dict, media_map: dict) -> None:
    """按消息类型回填 mediaFile，防跨分片 local_id 撞号错挂。

    导出管线只产出图片(3/47)与语音(34)两类媒体文件，媒体映射以
    (md5, localId, ts) 组合键（语音为 (\\x00voice, localId, ts)）为索引；
    多分片会话中各分片的 local_id 独立编号会撞号——图片与图片之间的互挂
    由 md5 组合键根治。文本/链接/系统消息不在类型门禁内，永不获媒体。
    图片槽位与语音槽位互斥，防图片/语音之间串型。

    合并转发（type 49）不产顶层媒体，但子项可能带图片——先回填子项，
    再走顶层类型门禁（issue #31）。
    """
    rec = msg.get("record")
    if rec:
        _attach_record_media(rec, media_map)
    t = msg.get("localType")
    if t not in (3, 47, 34):
        return
    ts = msg.get("createTime") or 0
    if t == 34:
        mf = media_map.get(_voice_key(msg.get("localId"), ts))
    else:
        mf = media_map.get(_media_key(msg.get("md5"), msg.get("bubbleMd5"),
                                      msg.get("localId"), ts))
    if not mf:
        # 审计 §4.1：媒体未回填正是"图丢了"的发生点，量=失败集合，可控。
        # ring=False：逐条埋点只进文件轨，防大导出冲掉日志页早期关键记录
        logger.detailed(
            "media",
            f"媒体未回填 localType={t} local_id={msg.get('localId')} ts={ts} "
            f"md5={(msg.get('md5') or '')[:8]} bm={(msg.get('bubbleMd5') or '')[:8]}",
            ring=False)
        return
    if (t == 34) != ("/voice_" in mf):
        logger.detailed(
            "media",
            f"媒体槽位不匹配 localType={t} local_id={msg.get('localId')} ts={ts} "
            f"mediaFile={mf}",
            ring=False)
        return
    msg["mediaFile"] = mf


def _log_media_failures(kind: str, failures: list, total: int) -> None:
    """聚合记录媒体失败（原因分布 + 逐条定位明细），终结导出失败全静默的盲区。

    ``failures`` 各项为 (md5, local_id, ts, reason) 元组——定位四件套在这里补齐
    （审计 §4.1：图片三级来源/Bubble 全靠 md5+local_id+ts，缺了无法对应到日志）。"""
    if not failures:
        return
    reasons = {}
    for _md5, _lid, _ts, r in failures:
        reasons[r] = reasons.get(r, 0) + 1
    top = "; ".join(f"{k} ×{v}" for k, v in
                    sorted(reasons.items(), key=lambda kv: -kv[1])[:3])
    logger.warn("export",
                f"[export] {kind}处理失败 {len(failures)}/{total}: {top}")
    for md5, local_id, ts, reason in failures:
        logger.detailed("export",
                        f"{kind}失败 local_id={local_id} ts={ts} "
                        f"md5={(md5 or '')[:8]} {reason}",
                        ring=False)


def _decrypt_media_parallel(acc_out_dir, account, chat, images, dest, progress=None,
                            should_cancel=None):
    """并行解密媒体图片。CPU 密集型 AES → 多进程池。

    images 每项是 (md5, bubble_md5, local_id, ts, msg_chat)——msg_chat 是该
    媒体所属的会话（顶层消息就是 chat；合并转发子项可能是原始消息所在的
    另一个会话，issue #31），attach/Bubble/Thumb 三级来源都按它定位。
    should_cancel：协作式取消探针，在结果消费循环里逐批检查。
    """
    dest.mkdir(parents=True, exist_ok=True)
    if not images:
        return {}

    # 任务元组：(acc_dir, account, chat, md5, bubble_md5, local_id, ts, dst)
    # chat / ts 必须带上：media.get_image() 的 attach 原图目录、Bubble 气泡缓存、
    # Thumb 缩略图三级来源都依赖它们，缺了就只剩 hardlink 一条路，大量图片解不出。
    tasks = []
    for i, (md5, bubble_md5, local_id, ts, msg_chat) in enumerate(images):
        fn = f"{i:04d}_{(md5 or 'img')[:12]}.jpg"
        dst = dest / fn
        tasks.append((str(acc_out_dir), account, msg_chat or chat, md5, bubble_md5,
                      local_id, ts, str(dst)))

    # 多进程并行解密
    n = min(os.cpu_count() or 4, len(tasks), 8)
    if n <= 1:
        # 串行兜底
        return _decrypt_media_serial(acc_out_dir, account, chat, images, dest,
                                     progress, should_cancel=should_cancel)

    from multiprocessing import Pool

    media_map = {}
    failures = []
    with Pool(n) as pool:
        for i, result in enumerate(pool.imap_unordered(_decrypt_one, tasks)):
            if should_cancel is not None and should_cancel():
                pool.terminate()
                raise ExportCancelled("媒体解密已取消")
            key, rel_path, ok, reason = result
            if ok:
                media_map[key] = rel_path
            else:
                # key = _media_key(md5, bubble_md5, local_id, ts)，自带定位信息；
                # 保留 (md5, local_id, ts, reason) 而非只留 reason（审计 B3 修正）
                failures.append((key[0], key[1], key[2], reason))
            if (i + 1) % 10 == 0 and progress:
                progress(0, f"媒体 {i + 1}/{len(images)}")
    _log_media_failures("图片", failures, len(images))
    return media_map


def _decrypt_media_serial(acc_out_dir, account, chat, images, dest, progress=None,
                          should_cancel=None):
    """串行解密（单核兜底）。"""
    media_map = {}
    failures = []
    for i, (md5, bubble_md5, local_id, ts, msg_chat) in enumerate(images):
        if should_cancel is not None and should_cancel():
            raise ExportCancelled("媒体解密已取消")
        fn = f"{i:04d}_{(md5 or 'img')[:12]}.jpg"
        dst = dest / fn
        out, reason = _try_decrypt(acc_out_dir, account, msg_chat or chat, md5,
                                   bubble_md5, local_id, ts, dst)
        if out:
            media_map[_media_key(md5, bubble_md5, local_id, ts)] = f"media/{out.name}"
        else:
            failures.append((md5 or bubble_md5 or "", local_id, ts, reason))
        if (i + 1) % 10 == 0 and progress:
            progress(0, f"媒体 {i + 1}/{len(images)}")
    _log_media_failures("图片", failures, len(images))
    return media_map


def _decrypt_one(task):
    """单张图片解密（子进程入口）。返回 (组合键, 相对路径, 是否成功, 失败原因)。"""
    acc_dir, account, chat, md5, bubble_md5, local_id, ts, dst = task
    out, reason = _try_decrypt(acc_dir, account, chat, md5, bubble_md5, local_id,
                               ts, Path(dst))
    return (_media_key(md5, bubble_md5, local_id, ts),
            f"media/{out.name}" if out else "", out is not None, reason)


def _try_decrypt(acc_dir, account, chat, md5, bubble_md5, local_id, ts, dst):
    """尝试解密单张图片。成功返回 (实际写出的 Path, "")，失败返回 (None, 原因)。

    Bug 修复：实际扩展名由图片内容决定（png/gif/jpg），必须把改写后的路径返回给
    调用方。原实现只返回 True/False，调用方却拿传入的 `.jpg` 占位名去拼 media
    引用，导致 PNG/GIF 图片在导出结果里指向不存在的文件。

    审计 §4.1：get_image 失败时第二返回值是细粒度 last_err（attach 解密失败 /
    V2 密钥未命中 / Bubble 未知格式 / 本地无原图），直接透传，不再统一覆盖成
    "未找到源文件或解密为空"——否则"没这个文件"与"有文件解不开"分不清。
    """
    try:
        body, info = media.get_image(account, md5, Path(acc_dir),
                                     chat=chat or None,
                                     local_id=local_id or None,
                                     ts=ts or None,
                                     bubble_md5=bubble_md5 or None)
        if body:
            # PR #28 同族缺陷：按子串猜扩展名会把 WebP 写成 .jpg。
            # info 是 mimetype（media._emit 产出 image/<ext>），直接取子类型；
            # 无法识别的类型显式失败，不静默落成 .jpg。
            if not info or not info.startswith("image/"):
                return None, f"未知媒体类型 {info!r}，不落盘"
            dst = Path(dst).with_suffix(f".{info.rsplit('/', 1)[-1]}")
            dst.write_bytes(body)
            return dst, ""
        return None, info or "未找到源文件或解密为空"
    except Exception as e:
        return None, f"{type(e).__name__}: {e}"


def _export_voice_media(acc_out_dir, account, chat, voices, dest, progress=None):
    """导出语音媒体：优先 WAV，缺少解码器时保留 SILK。"""
    dest.mkdir(parents=True, exist_ok=True)
    media_map = {}
    failures = []
    for i, (local_id, svr_id, ts) in enumerate(voices):
        out, reason = _try_export_voice(acc_out_dir, chat, local_id, svr_id, ts,
                                        dest / f"voice_{i:04d}_{local_id or 'msg'}")
        if out:
            media_map[_voice_key(local_id, ts)] = f"media/{out.name}"
        else:
            # 语音无 md5，首元组位留空（_log_media_failures 的定位口径统一）
            failures.append(("", local_id, ts, reason))
        if (i + 1) % 10 == 0 and progress:
            progress(0, f"语音 {i + 1}/{len(voices)}")
    _log_media_failures("语音", failures, len(voices))
    return media_map


def _try_export_voice(acc_dir, chat, local_id, svr_id, ts, dst_base):
    """读取并尝试转码一条语音。成功返回 (实际写出的 Path, "")，失败返回 (None, 原因)。

    转码失败不算失败：无本地解码器时保留清理后的 SILK 原文，供用户后续转换。"""
    try:
        body, info = voice.get_voice(Path(acc_dir), chat=chat or "",
                                     local_id=local_id or 0,
                                     svr_id=int(svr_id or 0), ts=ts or 0)
        if body is None:
            return None, "语音数据未找到"
        wav, meta = voice.transcode_voice(body, "wav")
        if wav is not None:
            dst = Path(dst_base).with_suffix(".wav")
            dst.write_bytes(wav)
            return dst, ""
        # 无本地解码器时不阻塞导出，保留清理后的 SILK 原文，供用户后续转换。
        dst = Path(dst_base).with_suffix(".silk")
        dst.write_bytes(body)
        return dst, ""
    except Exception as e:
        return None, f"{type(e).__name__}: {e}"


# 本地 Emoticon 缓存是微信私有封装解不开（api_chat.py /media/sticker 同结论），
# 本地图片管线解不出 47 消息时按消息 XML 的 cdnurl 下载明文表情落盘。
# 校验口径与 /media/sticker 一致：域名白名单 + 大小上限 + 图像签名。
_STICKER_MAX_BYTES = 10 * 1024 * 1024


def _export_sticker_media(dest, stickers, media_map, progress=None):
    """CDN 下载兜底导出动画表情，写好的映射并入 media_map。"""
    dest.mkdir(parents=True, exist_ok=True)
    todo = [(md5, bm, lid, ts, url) for md5, bm, lid, ts, url in stickers
            if url and sns_cdn.is_wechat_cdn(url)
            and _media_key(md5, bm, lid, ts) not in media_map]
    if not todo:
        return 0
    failures = []
    ok = 0
    for i, (md5, bm, lid, ts, url) in enumerate(todo):
        try:
            r = sns_cdn.fetch_emoji({"url": url}, timeout=15.0)
        except Exception as e:  # noqa: BLE001
            r = {"ok": False, "data": None, "ext": None,
                 "error": f"{type(e).__name__}: {e}"}
        if r.get("ok") and r.get("data") and len(r["data"]) <= _STICKER_MAX_BYTES:
            dst = dest / f"emoji_{i:04d}_{(md5 or 'stk')[:12]}.{r['ext']}"
            try:
                dst.write_bytes(r["data"])
                media_map[_media_key(md5, bm, lid, ts)] = f"media/{dst.name}"
                ok += 1
            except OSError as e:
                failures.append((md5 or bm or "", lid, ts,
                                 f"{type(e).__name__}: {e}"))
        else:
            failures.append((md5 or bm or "", lid, ts,
                             r.get("error") or "表情下载失败"))
        if (i + 1) % 10 == 0 and progress:
            progress(0, f"表情 {i + 1}/{len(todo)}")
    _log_media_failures("表情", failures, len(todo))
    return ok


class _ExportStats:
    """导出完整性统计（审计 G2）：解析成功/兜底计数、媒体回填计数。

    此前解析 fallback（[链接]/quote=None）与媒体失败全靠内部日志，用户拿到
    三万条导出后无从知道缺了什么；计数随导出写入 manifest.json。"""

    def __init__(self):
        self.type_counts = {}
        self.quote_ok = 0
        self.quote_fail = 0
        self.link_ok = 0
        self.record_count = 0
        self.record_items = 0
        self.fallback_labels = 0
        self.media_attached = 0
        self.truncated = 0      # 本格式实际截断的消息条数（P3-8）

    def observe(self, msg: dict) -> None:
        t = msg.get("localType")
        self.type_counts[t] = self.type_counts.get(t, 0) + 1
        if t == 57 or (t == 49 and has_refermsg(msg.get("rawContent") or "")):
            if msg.get("quote"):
                self.quote_ok += 1
            else:
                # 正文只记 hex 头 + 长度（_FILE_LOG 存未脱敏原文，正文截断落盘=明文落盘）
                raw = msg.get("rawContent") or ""
                logger.detailed(
                    "export",
                    f"quote解析失败 local_id={msg.get('localId')} "
                    f"ts={msg.get('createTime')} len={len(raw)} "
                    f"head={raw.encode('utf-8', 'replace')[:16].hex()}",
                    ring=False)
                self.quote_fail += 1
        elif msg.get("link"):
            self.link_ok += 1
        rec = msg.get("record")
        if rec:
            self.record_count += 1
            self.record_items += rec.get("count", 0)
        content = msg.get("content") or ""
        if content in _FALLBACK_LABELS_EXACT or _FALLBACK_TYPE_RE.match(content):
            logger.detailed(
                "export",
                f"兜底文案 local_id={msg.get('localId')} ts={msg.get('createTime')} "
                f"localType={t} len={len(content)} head={content[:32]!r}",
                ring=False)
            self.fallback_labels += 1
        if msg.get("mediaFile"):
            self.media_attached += 1

    def to_dict(self) -> dict:
        return {
            "type_counts": self.type_counts,
            "quote_parsed": self.quote_ok,
            "quote_failed": self.quote_fail,
            "link_parsed": self.link_ok,
            "merged_forward_messages": self.record_count,
            "merged_forward_items": self.record_items,
            "fallback_labels": self.fallback_labels,
            "media_attached": self.media_attached,
        }


def run_export(acc_out_dir: Path, account: str, chat: str, display: str,
               fmt: "str | list[str]", start_ts=None, end_ts=None,
               want_messages=True, want_media=True, want_voice=False,
               want_avatars=True,
               export_root: Path = None, pack: str = "zip",
               folder_name: str = None,
               template: str | None = None,
               progress=lambda pct, msg: None,
               should_cancel=None) -> dict:
    t0 = time.time()
    if should_cancel is not None and should_cancel():
        raise ExportCancelled("导出已取消")
    progress(1, f"开始导出: 账号={account}, 会话={chat}, 格式={fmt}")
    logger.detailed(
        "export",
        f"{logger.job_prefix()}开始导出 chat={chat} fmt={fmt} "
        f"range=[{start_ts},{end_ts}] "
        f"messages={want_messages} media={want_media} voice={want_voice} "
        f"avatars={want_avatars}")

    # 联系人缓存：全流程只加载一次，供两遍扫描共用（避免重复读 contact.db）
    names = _contact_names(acc_out_dir)

    # ── 第一遍：轻量采集元数据 ─────────────────────────────
    progress(3, "扫描消息元数据…")
    count, first_ts, last_ts, senders, images, voices, stickers = _collect_metadata(
        acc_out_dir, chat, start_ts, end_ts, account, names)
    # P2-1/P2-3：单遍逻辑条数——want_messages=False（仅元数据/媒体）时为 0。
    # 多格式导出的 message_count/manifest written_count 都用它，不按格式数累加。
    output_count = count if want_messages else 0
    progress(10, f"共 {count} 条消息，{len(images)} 张图片，{len(voices)} 条语音")
    logger.detailed(
        "export",
        f"扫描: msgs={count} images={len(images)} voices={len(voices)} "
        f"stickers={len(stickers)} "
        f"senders={len(senders)} range=[{first_ts},{last_ts}]")

    display = display or names.get(chat, chat) or chat
    safe = _safe_name(display)
    # P3：时间戳到毫秒——秒级精度下同一会话一秒内两次导出（批量/重试场景）
    # 会产生同名 zip，后者 make_archive 静默覆盖前者
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:-3]
    root = (export_root or _paths.exports_root())
    export_dir = root / _safe_name(folder_name or f"{stamp}_{safe}")
    export_dir.mkdir(parents=True, exist_ok=True)

    session = {
        "wxid": chat, "nickname": names.get(chat, chat),
        "displayName": display,
        "type": "群聊" if chat.endswith("@chatroom") else "私聊",
        "platform": "wechat", "isGroup": chat.endswith("@chatroom"),
        "ownerId": account,
        "firstTimestamp": first_ts, "lastTimestamp": last_ts,
        # P2-3：仅元数据模式不产出消息，会话头计数与之保持一致
        "messageCount": output_count,
    }

    # ── 头像 ──────────────────────────────────────────────
    avatar_map = {}
    stats_ava = 0
    if want_avatars and count > 0:
        if should_cancel is not None and should_cancel():
            raise ExportCancelled("导出已取消")
        progress(15, "提取头像…")
        dest = export_dir / "avatars"
        dest.mkdir(exist_ok=True)
        avatar_map = collect_avatars(acc_out_dir, list(senders), dest, progress)
        stats_ava = len(avatar_map)
        progress(30, f"头像 {stats_ava}/{len(senders)}")

    # ── 媒体（图片并行解密，语音优先转 WAV，表情 CDN 下载兜底）──────────
    stats_media = 0
    stats_voice = 0
    stats_sticker = 0
    media_map = {}
    if (want_media and (images or stickers)) or (want_voice and voices):
        if should_cancel is not None and should_cancel():
            raise ExportCancelled("导出已取消")
        media_dest = export_dir / "media"
        img_total = len(images) if want_media else 0
        voice_total = len(voices) if want_voice else 0
        if want_media and images:
            progress(35, f"解密图片（{len(images)} 张）…")
            image_map = _decrypt_media_parallel(acc_out_dir, account, chat, images,
                                                media_dest, progress,
                                                should_cancel=should_cancel)
            media_map.update(image_map)
            stats_media = len(image_map)
        if want_voice and voices:
            if should_cancel is not None and should_cancel():
                raise ExportCancelled("导出已取消")
            progress(55, f"转码语音（{len(voices)} 条）…")
            voice_map = _export_voice_media(acc_out_dir, account, chat, voices,
                                            media_dest, progress)
            media_map.update(voice_map)
            stats_voice = len(voice_map)
        if want_media and stickers:
            if should_cancel is not None and should_cancel():
                raise ExportCancelled("导出已取消")
            progress(65, f"下载表情（{len(stickers)} 个）…")
            stats_sticker = _export_sticker_media(media_dest, stickers,
                                                  media_map, progress)
        progress(70, f"媒体处理完成: 图片 {stats_media}/{img_total}，"
                     f"语音 {stats_voice}/{voice_total}，"
                     f"表情 {stats_sticker}/{len(stickers) if want_media else 0}")
        logger.detailed(
            "export",
            f"媒体完成 img={stats_media}/{img_total} "
            f"voice={stats_voice}/{voice_total} "
            f"sticker={stats_sticker}/{len(stickers) if want_media else 0}")

    # ── 第二遍：流式写出（fmt 可传列表，一次写出多种格式）──
    # 扫描/头像/媒体只做一次；每种格式各自重放 message_stream（DB 重查
    # 毫秒级）并各自计数，避免 manifest 的解析统计被乘以格式数。
    # 计数口径（P2-1）：message_count = 单遍逻辑条数 output_count；
    # written_total = 多格式实际写入总量（仅进度/对账/排查用）。
    fmts = [str(f).lower() for f in fmt] if isinstance(fmt, (list, tuple)) \
        else [str(fmt).lower()]
    fname = f"{safe}_{stamp}"
    files: list[str] = []
    written_total = 0
    # P3-8：每格式实际截断条数（无截断也记 0，供 manifest/界面统一读取）
    truncated_by_fmt: dict = {}
    stats_manifest = None
    for fi, one_fmt in enumerate(fmts):
        if should_cancel is not None and should_cancel():
            raise ExportCancelled("导出已取消")
        progress(75 + int(fi * 12 / len(fmts)),
                 f"写入文件（{one_fmt}，{fi + 1}/{len(fmts)}）…")
        plugin_writer = _plugin_export_format(one_fmt)
        ext = {"json": "json", "html": "html", "txt": "txt", "csv": "csv",
               "markdown": "md", "toml": "toml", "sqlite": "db",
               "xlsx": "xlsx"}.get(one_fmt)
        if ext is None:
            # 插件格式：用插件声明的 ext，兜底回退 fmt 本身
            ext = (plugin_writer.ext if plugin_writer is not None else None) or one_fmt
        out_file = export_dir / f"{fname}.{ext}"
        stats = _ExportStats()
        if stats_manifest is None:
            stats_manifest = stats

        def export_stream_with_media(_stats=stats):
            # P2-3：want_messages=False（仅元数据/媒体）时不产出任何消息记录，
            # 生成器直接返回——JSON/TXT/CSV/Markdown/插件都拿到空序列，
            # stats.observe 随之跳过（统计只在真正写消息时有意义）。
            if not want_messages:
                return
            for msg in message_stream(acc_out_dir, chat, start_ts, end_ts, account, names):
                _attach_media(msg, media_map)
                _stats.observe(msg)
                yield msg

        # ── 传入缓存的 names，避免重复加载 ───────────────────
        if plugin_writer is not None:
            stream = export_stream_with_media() if want_messages else iter(())
            written = _run_plugin_writer(plugin_writer, out_file, session,
                                         stream, media_map, progress,
                                         acc_out_dir, chat, start_ts, end_ts, account, names)
        elif one_fmt == "json":
            written = stream_export_json(out_file, session, export_stream_with_media(), progress)
        elif one_fmt == "html":
            written = _write_html_streaming(out_file, acc_out_dir, chat, start_ts,
                                            end_ts, account, names, media_map,
                                            avatar_map, progress, display, stats,
                                            session=session, senders=senders,
                                            template=template,
                                            media_enabled=want_media,
                                            want_messages=want_messages)
        elif one_fmt == "txt":
            written = stream_export_txt(out_file, session, export_stream_with_media(), progress)
        elif one_fmt == "csv":
            written = stream_export_csv(out_file, export_stream_with_media(), progress)
        elif one_fmt == "markdown":
            written = stream_export_md(out_file, session, export_stream_with_media(), progress)
        elif one_fmt in ("toml", "sqlite", "xlsx"):
            # 这些格式需要全量数据 → 降级为流式分批（每批 500 条）
            written = _write_batch(out_file, one_fmt, session,
                                   acc_out_dir, chat, start_ts, end_ts, account, names, media_map, progress,
                                   stats=stats, want_messages=want_messages)
        else:
            raise ValueError(f"未知格式: {one_fmt}")
        files.append(str(out_file))
        written_total += written
        truncated_by_fmt[one_fmt] = stats.truncated
        progress(87, f"写入完成（{one_fmt}）: {written} 条")
        logger.detailed("export",
                        f"{logger.job_prefix()}写出完成 fmt={one_fmt} "
                        f"written={written} collected={count} "
                        f"messages={want_messages}")
        # 对账按单格式 written 比单遍采集数；want_messages=False 时 0 条是
        # 预期行为，不误报为缺消息
        if want_messages and count and written != count:
            logger.warn("export",
                        f"[export] 条数对账不一致 (fmt={one_fmt}): 采集 {count} 条 vs 实际写出 "
                        f"{written} 条（可能有分片读取失败、消息缺失，详见上方日志）")

    # ── 导出完整性清单（审计 G2）──────────────────────────
    # 各类型解析成功/兜底计数、媒体失败分布随导出落盘，用户能直接看到
    # "这次导出缺了什么"，而不是事后靠翻日志。
    export_manifest = {
        "generator": GENERATOR, "export_version": EXPORT_VERSION,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "chat": chat, "display": display,
        # P2-1：written_count 是单遍逻辑条数（多格式不翻倍）；
        # written_total 是多格式实际写入总量（对账/排查用）
        "collected_count": count, "written_count": output_count,
        "written_total": written_total,
        "truncated": truncated_by_fmt,
        "files": files,
        "parse": stats_manifest.to_dict(),
        "media": {
            # 口径修正（P3）：未请求图片时 requested 应为 0，而非实际采集数
            "images_requested": len(images) if want_media else 0,
            "images_ok": stats_media,
            "voices_requested": len(voices) if want_voice else 0,
            "voices_ok": stats_voice,
            "stickers_requested": len(stickers) if want_media else 0,
            "stickers_ok": stats_sticker,
        },
        "avatar_count": stats_ava,
    }
    try:
        (export_dir / "manifest.json").write_text(
            json.dumps(export_manifest, ensure_ascii=False, indent=1),
            encoding="utf-8")
        logger.info("export",
                    f"[export] manifest.json 已生成: 引用解析 {stats_manifest.quote_ok}"
                    f"{'/失败 ' + str(stats_manifest.quote_fail) if stats_manifest.quote_fail else ''}, "
                    f"链接 {stats_manifest.link_ok}, 合并转发 {stats_manifest.record_count}"
                    f"（{stats_manifest.record_items} 条子消息）, 兜底文案 {stats_manifest.fallback_labels}")
        logger.detailed("export", f"type_counts={stats_manifest.type_counts}")
    except OSError as e:
        logger.warn("export", f"[export] manifest.json 写入失败: {e}")

    # ── 插件：导出后处理（before_zip）─────────────────────
    # 必须在打包 zip **之前**运行：zip 一旦生成会 rmtree 掉 export_dir，
    # after_zip 阶段只剩 zip 文件本身可操作。
    export_result = {
        "export_dir": str(export_dir),
        "file": files[0] if files else None,
        "files": list(files),
        "format": fmt, "pack": pack,
        "message_count": output_count,
        "written_total": written_total,
    }
    _run_after_export("before_zip", {
        **export_result, "account": account, "chat": chat,
        "display": display, "names": names, "media_map": media_map,
        "progress": progress,
    })

    # ── 打包 ─────────────────────────────────────────────
    zip_path = None
    file_path = files[0] if files else None
    files_out = list(files)
    if pack == "zip":
        progress(92, "打包 zip…")
        zip_path = shutil.make_archive(str(root / f"{fname}_{'_'.join(fmts)}"), "zip",
                                       root_dir=export_dir)
        shutil.rmtree(export_dir, ignore_errors=True)
        # 目录已删，file 若仍指向目录内路径就是死链：「每会话一个 ZIP」
        # （pack="each" → 每会话 run_export(pack="zip")）模式下前端拿它
        # 渲染"下载文件"链接，点击必然 404。置空，前端对空值不渲染链接。
        file_path = None
        files_out = []
        export_result["file"] = None
        export_result["files"] = []
        # ── 插件：导出后处理（after_zip）───────────────────
        # 此时 export_dir 已被删除，只提供 zip 路径。
        _run_after_export("after_zip", {
            **export_result, "export_dir": None, "zip": zip_path,
            "account": account, "chat": chat, "display": display,
            "progress": progress,
        })

    progress(100, "导出完成")
    return {
        "export_dir": str(root) if pack == "zip" else str(export_dir),
        "zip": zip_path,
        "file": file_path,
        "files": files_out,
        "format": fmt, "pack": pack,
        "message_count": output_count, "written_total": written_total,
        "truncated": truncated_by_fmt,
        "media_count": stats_media + stats_voice,
        "image_count": stats_media, "voice_count": stats_voice,
        "avatar_count": stats_ava,
        "duration_ms": int((time.time() - t0) * 1000),
    }


# ── 插件桥接（无插件时零开销）────────────────────────────────

def _plugin_export_format(fmt: str):
    """按 fmt 取插件声明的导出格式（内置格式优先，插件只补新格式）。"""
    try:
        from siwx.plugins import ensure_loaded, registry
        ensure_loaded()
        return registry.find_export_format(fmt)
    except Exception as e:
        logger.warn("plugin", f"导出格式查询失败: {e}")
        return None


def _run_plugin_writer(w, path, session, stream, media_map, progress,
                       acc_out_dir, chat, start_ts, end_ts, account, names) -> int:
    """调用插件导出写入器。

    插件契约：`writer(path, ctx) -> int`（返回写入条数）。

    ctx 关键字段
    ------------
    stream : 惰性消息生成器（**导出形状**，非前端形状），每项字段：
             localId / platformMessageId / createTime / localType / typeName /
             rawContent / content / isSend / senderUsername / senderDisplayName /
             md5 / bubbleMd5 / voice / quote / link / mediaFile
             ``content`` 已按类型格式化（图片为 "[图片]" 等）。
    session / names / media_map / account / chat / start_ts / end_ts
    out_dir / progress / fmt / ext

    media_map 键为 (md5, localId, ts) 组合键（语音为 (\\x00voice, localId, ts)，
    见 _media_key/_voice_key）——旧版裸 localId 键在多分片会话下会跨分片撞号
    互挂图片。stream 只能迭代一次；需要多次遍历请自行 list() 缓存。
    """
    from siwx import logger as log
    plugin = w.meta.name if w.meta else "?"
    ctx = {
        "session": session, "stream": stream, "media_map": media_map,
        "account": account, "chat": chat, "start_ts": start_ts,
        "end_ts": end_ts, "names": names, "out_dir": str(acc_out_dir),
        "progress": progress, "fmt": w.fmt, "ext": w.ext,
    }
    try:
        n = w.writer(path, ctx)
        return int(n) if isinstance(n, (int, float)) else 0
    except Exception as e:
        log.error("plugin", f"{plugin} 导出格式 {w.fmt} 写入失败: {e}")
        raise


def _run_after_export(when: str, ctx: dict) -> None:
    """运行 when 阶段的插件后处理钩子（逐插件隔离，不阻塞导出）。"""
    try:
        from siwx.plugins import ensure_loaded, registry
        ensure_loaded()
        hooks = registry.after_export
    except Exception as e:
        logger.warn("plugin", f"after_export 钩子加载失败: {e}")
        return
    if not hooks:
        return
    from siwx import logger as log
    for _i, h in hooks.sorted_items():
        if (h.when or "before_zip") != when:
            continue
        plugin = h.meta.name if h.meta else (h.name or "?")
        try:
            h.run(dict(ctx))
        except Exception as e:
            log.warn("plugin", f"{plugin}.after_export({when}) 失败: {e}")


def _write_html_streaming(path, acc_dir, chat, start_ts, end_ts, account,
                          names, media_map, avatar_map, progress=None,
                          display=None, stats=None, session=None, senders=None,
                          template=None, media_enabled=True, want_messages=True):
    """HTML 真流式导出（P1 修复）：消息边读边写文件句柄，内存 O(1)。

    旧实现把全部消息 accumulate 进 lines 再一次性 render_html——3 万条
    会话 = 全量消息 dict + 单个巨型 HTML 字符串，与模块"万条不 OOM"目标
    矛盾。成员表由第一遍扫描的 senders 集合预构建（meta.messageCount 用
    第一遍计数，尾部 MSG_COUNT 用实际写出数，二者以 run_export 的对账
    warn 兜底差异）。

    want_messages=False（P2-3，仅元数据/媒体）：跳过消息循环，但仍写
    头部、尾部与成员表，得到可打开的空消息页（MSG_COUNT=0）。
    """
    from siwx.html_template import (
        stream_html_head, stream_html_msg, stream_html_tail,
    )
    from siwx import wx_faces, wx_maps
    if session is None:
        session = {"wxid": chat,
                   "displayName": display or names.get(chat, chat) or chat,
                   "isGroup": chat.endswith("@chatroom"),
                   "firstTimestamp": 0, "lastTimestamp": 0,
                   "ownerId": account, "messageCount": 0}
    members = [{"id": un, "name": names.get(un, un) or un,
                "avatar": avatar_map.get(un, "")}
               for un in sorted(senders or ())]

    count = 0
    faces_used: set = set()
    map_keys: set = set()
    with open(path, "w", encoding="utf-8") as f:
        stream_html_head(f, session, members, avatar_map, template=template)
        # P2-3：仅元数据模式不写消息，头/尾/成员表照常输出
        msgs = (message_stream(acc_dir, chat, start_ts, end_ts, account, names)
                if want_messages else ())
        for msg in msgs:
            _attach_media(msg, media_map)
            if stats:
                stats.observe(msg)
            faces_used |= wx_faces.used_from_message(msg)
            map_keys |= wx_maps.used_from_message(msg)
            if stats and len(msg.get("rawContent") or "") > HTML_RAW_CONTENT_LIMIT:
                stats.truncated += 1
            stream_html_msg(f, msg, first=(count == 0))
            count += 1
            if count % 500 == 0:
                f.flush()
                if progress:
                    progress(0, f"已写出 {count} 条…")
        stream_html_tail(f, count, session["displayName"],
                         faces=wx_faces.datauris(faces_used),
                         maps=wx_maps.datauris(map_keys),
                         template=template,
                         media_enabled=media_enabled)
    if stats and stats.truncated:
        logger.detailed("html",
                        f"HTML rawContent截断(>{HTML_RAW_CONTENT_LIMIT}字) "
                        f"{stats.truncated}/{count} 条", ring=False)

    # 审计 §4.6：服务端统计埋点落在调用方（html_template 保持零业务依赖）
    logger.detailed(
        "html",
        f"HTML写出: msgs={count} media={stats.media_attached if stats else 0} "
        f"quote_fail={stats.quote_fail if stats else 0} "
        f"fallback_labels={stats.fallback_labels if stats else 0} "
        f"type_counts={stats.type_counts if stats else {}}")
    return count


def _write_batch(path, fmt, session, acc_dir, chat, start_ts, end_ts, account, names,
                 media_map, progress=None, stats=None, want_messages=True):
    """TOML/SQLite/XLSX 流式分批写入。

    want_messages=False（P2-3）：只写文件头/表结构，不写任何消息，返回 0。"""
    if fmt == "toml":
        return _write_toml_batch(
            path, session, acc_dir, chat, start_ts, end_ts, account, names,
            media_map, progress, stats, want_messages)
    elif fmt == "sqlite":
        return _write_sqlite_batch(
            path, session, acc_dir, chat, start_ts, end_ts, account, names,
            media_map, progress, stats, want_messages)
    elif fmt == "xlsx":
        return _write_xlsx_batch(
            path, acc_dir, chat, start_ts, end_ts, account, names,
            media_map, progress, stats, want_messages)


def _toml_value(v) -> str:
    """TOML 标量序列化。

    TOML 只接受小写 ``true``/``false``，Python 的 ``str(True)`` 会拼出
    ``isGroup = True`` 让 ``tomllib`` 直接报错；数值原样输出；字符串走
    ``json.dumps``（JSON 字符串转义是 TOML basic string 的子集）；
    TOML 没有 null，None 退化为空字符串。调用方仍需跳过 dict 等复合值。
    """
    if v is None:
        return '""'
    if isinstance(v, bool):          # 必须在 int 之前：bool 是 int 的子类
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    return json.dumps(str(v), ensure_ascii=False)


def _write_toml_batch(path, session, acc_dir, chat, start_ts, end_ts, account, names, media_map, progress, stats=None, want_messages=True):
    header = ["[exportInfo]", f'version = {_toml_value("1.0")}',
              f'generator = {_toml_value(GENERATOR)}',
              f'exportedAt = {_toml_value(int(time.time()))}',
              "", "[session]"]
    for k, v in session.items():
        if isinstance(v, dict):
            pass
        else:
            header.append(f'{k} = {_toml_value(v)}')
    header.append("")
    count = 0
    # 逐条写盘：旧实现把整篇 messages 先 append 进 lines 列表、最后一次 write_text，
    # 峰值 ≈ 输出文件的 3 倍（lines + join 串 + encode 缓冲），最长会话约 12.4 万条
    # 时能到数百 MB；JSON/HTML 侧本就是 O(1) 流式，TOML 反而更差。
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(header))
        f.write("\n")
        # P2-3：仅元数据模式只写 [exportInfo]/[session] 头，不写 [[messages]]
        msgs = (message_stream(acc_dir, chat, start_ts, end_ts, account, names)
                if want_messages else ())
        for msg in msgs:
            _attach_media(msg, media_map)
            if stats:
                stats.observe(msg)
            count += 1
            block = ["[[messages]]",
                     f"localId = {_toml_value(msg['localId'])}",
                     f"createTime = {_toml_value(datetime.fromtimestamp(msg['createTime']).strftime('%Y-%m-%d %H:%M:%S'))}",
                     f"type = {_toml_value(msg['typeName'])}",
                     f"sender = {_toml_value(msg['senderDisplayName'])}",
                     f"isSend = {_toml_value(msg['isSend'])}",
                     f"content = {_toml_value(msg['content'])}"]
            if msg.get("mediaFile"):
                block.append(f"mediaFile = {_toml_value(msg['mediaFile'])}")
            block.append("")
            f.write("\n".join(block))
            f.write("\n")
            if count % 500 == 0 and progress:
                progress(0, f"已写入 {count} 条…")
    return count


def _write_sqlite_batch(path, session, acc_dir, chat, start_ts, end_ts, account, names, media_map, progress, stats=None, want_messages=True):
    if path.exists():
        path.unlink()
    conn = sqlite3.connect(path)
    conn.executescript("""
        CREATE TABLE session (wxid TEXT, nickname TEXT, type TEXT, isGroup INTEGER,
            messageCount INTEGER, firstTimestamp INTEGER, lastTimestamp INTEGER);
        CREATE TABLE messages (localId INTEGER, createTime INTEGER, formattedTime TEXT,
            localType INTEGER, typeName TEXT, isSend INTEGER, senderUsername TEXT,
            senderDisplayName TEXT, content TEXT, rawContent TEXT, mediaFile TEXT);
        CREATE INDEX msg_idx ON messages(createTime);
    """)
    conn.execute("INSERT INTO session VALUES (?,?,?,?,?,?,?)",
                 (session["wxid"], session.get("nickname", ""), session["type"],
                  session["isGroup"], session["messageCount"],
                  session["firstTimestamp"], session["lastTimestamp"]))
    count = 0
    # P2-3：仅元数据模式保留 session 表与 messages 表结构，不插任何消息行
    msgs = (message_stream(acc_dir, chat, start_ts, end_ts, account, names)
            if want_messages else ())
    for msg in msgs:
        _attach_media(msg, media_map)
        if stats:
            stats.observe(msg)
        count += 1
        raw = msg.get("rawContent") or ""
        if stats and len(raw) > SQLITE_RAW_CONTENT_LIMIT:
            stats.truncated += 1
        conn.execute("INSERT INTO messages VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                     (msg["localId"], msg["createTime"],
                      datetime.fromtimestamp(msg["createTime"]).strftime("%Y-%m-%d %H:%M:%S"),
                      msg["localType"], msg["typeName"], msg["isSend"],
                      msg["senderUsername"], msg["senderDisplayName"],
                      msg["content"], raw[:SQLITE_RAW_CONTENT_LIMIT],
                      msg.get("mediaFile")))
        if count % 1000 == 0:
            conn.commit()
            if progress:
                progress(0, f"已写入 {count} 条…")
    conn.commit()
    conn.close()
    if stats and stats.truncated:
        logger.detailed("export",
                        f"SQLite rawContent截断(>{SQLITE_RAW_CONTENT_LIMIT}字) "
                        f"{stats.truncated}/{count} 条", ring=False)
    return count


def _write_xlsx_batch(path, acc_dir, chat, start_ts, end_ts, account, names, media_map, progress, stats=None, want_messages=True):
    from openpyxl import Workbook
    # write_only=True：默认 Workbook() 会把每个 Cell 对象都留在内存里，长会话
    # （12.4 万条 × 7 列 = 近百万 Cell）占用远超文件本身；write_only 只保留
    # 写游标，内存 O(1)。read_only 读回/校验不受影响。
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("聊天记录")
    ws.append(["localId", "时间", "类型", "发送者", "是否自己", "内容", "媒体文件"])
    count = 0
    # P2-3：仅元数据模式保留表头行，不追加任何消息行
    msgs = (message_stream(acc_dir, chat, start_ts, end_ts, account, names)
            if want_messages else ())
    for msg in msgs:
        _attach_media(msg, media_map)
        if stats:
            stats.observe(msg)
        count += 1
        if stats and len(msg["content"]) > XLSX_CELL_LIMIT:
            stats.truncated += 1
        ws.append([msg["localId"],
                   datetime.fromtimestamp(msg["createTime"]).strftime("%Y-%m-%d %H:%M:%S"),
                   msg["typeName"], msg["senderDisplayName"],
                   "是" if msg["isSend"] else "否", msg["content"][:XLSX_CELL_LIMIT],
                   msg.get("mediaFile") or ""])
        if count % 1000 == 0 and progress:
            progress(0, f"已写入 {count} 条…")
    wb.save(path)
    if stats and stats.truncated:
        logger.detailed("export",
                        f"XLSX内容截断(>{XLSX_CELL_LIMIT}字) {stats.truncated}/{count} 条，"
                        f"换 JSON/HTML 格式可得全文", ring=False)
    return count


def run_export_multi(acc_out_dir: Path, account: str, chats: list,
                     fmt: "str | list[str]" = "json",
                     start_ts=None, end_ts=None,
                     want_messages=True, want_media=False, want_voice=False,
                     want_avatars=False,
                     export_root: Path = None, pack: str = "folder",
                     template: str | None = None,
                     progress=lambda pct, msg: None,
                     should_cancel=None) -> dict:
    """多会话批量导出。should_cancel：协作式取消探针——每个会话开始前、
    每个阶段检查点生效；取消是优雅停止，已完成会话的产物保留。"""
    t0 = time.time()
    n = max(len(chats), 1)
    # P3：毫秒级时间戳，防同秒批量导出目录名冲突（同因 as run_export）
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:-3]
    root = Path(export_root or _paths.exports_root())
    total_dir = root / f"export_{stamp}"
    total_dir.mkdir(parents=True, exist_ok=True)
    progress(1, f"共 {n} 个会话 · 输出目录 {total_dir.name}/ · 打包={pack}")

    results, ok_n = [], 0
    cancelled = False
    for i, item in enumerate(chats):
        if should_cancel is not None and should_cancel():
            cancelled = True
            progress(int(100 * i / n), f"已取消（完成 {i}/{n}）")
            break
        chat = item.get("chat") if isinstance(item, dict) else str(item)
        display = (item.get("display") or "") if isinstance(item, dict) else ""
        base, endp = int(100 * i / n), int(100 * (i + 1) / n)

        def sub(pct, msg, _b=base, _e=endp, _i=i):
            progress(_b + int(pct * (_e - _b) / 100), f"[{_i + 1}/{n}] {msg}")

        try:
            res = run_export(acc_out_dir, account, chat, display, fmt,
                             start_ts, end_ts,
                             want_messages=want_messages,
                             want_media=want_media,
                             want_voice=want_voice,
                             want_avatars=want_avatars,
                             export_root=total_dir,
                             folder_name=f"{i + 1:02d}_{display or chat}",
                             pack=("zip" if pack == "each" else "none"),
                             template=template,
                             progress=sub,
                             should_cancel=should_cancel)
            ok_n += 1
            files_out = list(res.get("files") or [])
            if res.get("file"):
                files_out.append(res["file"])
            # 产物总大小：报告里给个数，前端就能标"哪个会话拖慢/撑大了导出"
            size_bytes = 0
            for f in files_out:
                try:
                    size_bytes += Path(f).stat().st_size
                except OSError:
                    pass
            results.append({"chat": chat, "display": display or chat,
                            "message_count": res.get("message_count", 0),
                            "media_count": res.get("media_count", 0),
                            "image_count": res.get("image_count", 0),
                            "voice_count": res.get("voice_count", 0),
                            "avatar_count": res.get("avatar_count", 0),
                            "file": res.get("file"),
                            "files": res.get("files", []),
                            "size_bytes": size_bytes,
                            "duration_ms": res.get("duration_ms", 0)})
            progress(endp, f"[{i + 1}/{n}] ✔ {display or chat} ({res.get('message_count', 0)} 条)")
        except ExportCancelled:
            cancelled = True
            progress(endp, f"[{i + 1}/{n}] 已取消（会话 {display or chat} 未完成）")
            break
        except Exception as e:
            logger.error("export",
                         f"{logger.job_prefix()}会话导出失败 chat={chat}: "
                         f"{type(e).__name__}: {e}")
            results.append({"chat": chat, "display": display or chat, "error": str(e)})
            progress(endp, f"[{i + 1}/{n}] ✗ {display or chat}: {e}")

    zip_path, zips = None, []
    if not cancelled:
        if pack == "each":
            zips = sorted(str(p) for p in total_dir.glob("*.zip"))
            progress(97, f"每会话 zip 共 {len(zips)} 个")
        elif pack == "single":
            progress(96, "打包整体 zip…")
            zip_path = shutil.make_archive(str(root / total_dir.name), "zip",
                                           root_dir=total_dir)
            progress(98, f"zip 完成: {Path(zip_path).stat().st_size / 1048576:.1f} MB")

    progress(100, "已取消，已完成部分保留" if cancelled else "导出完成")
    zip_size = None
    if zip_path:
        try:
            zip_size = Path(zip_path).stat().st_size
        except OSError:
            zip_size = None
    return {
        "total_dir": str(total_dir), "zip": zip_path, "zips": zips,
        "zip_size": zip_size,
        "pack": pack, "format": fmt, "sessions": results, "ok_count": ok_n,
        "cancelled": cancelled,
        "message_count": sum(r.get("message_count", 0) for r in results),
        "media_count": sum(r.get("media_count", 0) for r in results),
        "image_count": sum(r.get("image_count", 0) for r in results),
        "voice_count": sum(r.get("voice_count", 0) for r in results),
        "avatar_count": sum(r.get("avatar_count", 0) for r in results),
        "duration_ms": int((time.time() - t0) * 1000),
    }
